"""YouTube URL, metadata, transcript, and structured Markdown note workflow."""

from __future__ import annotations

import asyncio
import json
import re
import tempfile
import urllib.request
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any, Callable, Iterable
from urllib.parse import parse_qs, urlencode, urlparse

from core import gemini

from .framework import WorkflowContext, WorkflowRunner, WorkflowStep


@dataclass(frozen=True, slots=True)
class YouTubeReference:
    """Normalized YouTube video reference."""

    video_id: str
    canonical_url: str
    playlist_id: str = ""


@dataclass(frozen=True, slots=True)
class TranscriptSegment:
    """Transcript text with its source timestamp."""

    start: float
    duration: float
    text: str


def parse_youtube_url(value: str) -> YouTubeReference:
    """Parse watch, short, embed, shorts, and playlist-item URLs."""
    parsed = urlparse(str(value or '').strip())
    host = parsed.netloc.casefold().split(':', 1)[0]
    if host.startswith('www.'):
        host = host[4:]
    if host not in {'youtube.com', 'm.youtube.com', 'music.youtube.com', 'youtu.be'}:
        raise ValueError('not a supported YouTube URL')
    query = parse_qs(parsed.query)
    video_id = query.get('v', [''])[0]
    if host == 'youtu.be' and not video_id:
        video_id = parsed.path.strip('/').split('/', 1)[0]
    if not video_id:
        parts = [part for part in parsed.path.split('/') if part]
        if parts and parts[0] in {'shorts', 'embed', 'v'} and len(parts) > 1:
            video_id = parts[1]
    if not re.fullmatch(r'[A-Za-z0-9_-]{11}', video_id or ''):
        raise ValueError('could not extract a valid YouTube video id')
    playlist_id = query.get('list', [''])[0]
    canonical_query = {'v': video_id}
    if playlist_id:
        canonical_query['list'] = playlist_id
    return YouTubeReference(
        video_id=video_id,
        canonical_url='https://www.youtube.com/watch?' + urlencode(canonical_query),
        playlist_id=playlist_id,
    )


def chunk_transcript(segments: Iterable[TranscriptSegment], max_chars: int = 12_000,
                     overlap_chars: int = 600) -> list[list[TranscriptSegment]]:
    """Split transcript segments by character budget while preserving overlap."""
    if max_chars <= 0 or overlap_chars < 0 or overlap_chars >= max_chars:
        raise ValueError('overlap must be non-negative and smaller than max_chars')
    source = list(segments)
    chunks: list[list[TranscriptSegment]] = []
    current: list[TranscriptSegment] = []
    size = 0
    for segment in source:
        segment_size = len(segment.text) + 1
        if current and size + segment_size > max_chars:
            chunks.append(current)
            carried = []
            carried_size = 0
            for previous in reversed(current):
                carried.insert(0, previous)
                carried_size += len(previous.text) + 1
                if carried_size >= overlap_chars:
                    break
            current = carried
            size = carried_size
        current.append(segment)
        size += segment_size
    if current:
        chunks.append(current)
    return chunks


def _safe_filename(title: str) -> str:
    """Return a filesystem-safe note stem."""
    cleaned = re.sub(r'[<>:"/\\|?*\x00-\x1f]', '-', title).strip(' .')
    return (cleaned[:120] or 'youtube-notes')


def _timestamp_link(reference: YouTubeReference, seconds: float) -> str:
    return f"https://youtu.be/{reference.video_id}?t={max(0, int(seconds))}"


class YouTubeNotesWorkflow:
    """Produce and save complete Markdown notes for one YouTube video."""

    def __init__(self, notes_dir: str | Path | None = None,
                 notify: Callable[[str], None] | None = None):
        self.notes_dir = Path(notes_dir or (Path.home() / 'Documents' / 'AURA Notes'))
        self.notify = notify or (lambda _message: None)

    def metadata(self, reference: YouTubeReference) -> dict[str, Any]:
        """Fetch metadata through yt-dlp's official extractor."""
        try:
            import yt_dlp
        except ImportError as exc:
            raise RuntimeError('yt-dlp is required for YouTube metadata') from exc
        options = {'quiet': True, 'no_warnings': True, 'skip_download': True, 'noplaylist': True}
        with yt_dlp.YoutubeDL(options) as downloader:
            info = downloader.extract_info(reference.canonical_url, download=False)
        return {
            'title': str(info.get('title') or reference.video_id),
            'channel': str(info.get('channel') or info.get('uploader') or 'Unknown'),
            'duration': int(info.get('duration') or 0),
            'publish_date': str(info.get('upload_date') or ''),
            'url': reference.canonical_url,
            'subtitles': info.get('subtitles') or {},
            'automatic_captions': info.get('automatic_captions') or {},
        }

    @staticmethod
    def _official_transcript(video_id: str) -> list[TranscriptSegment]:
        from youtube_transcript_api import YouTubeTranscriptApi
        transcript = YouTubeTranscriptApi().fetch(video_id, languages=['en', 'tr', 'de', 'fr', 'es'])
        return [TranscriptSegment(float(item.start), float(item.duration), str(item.text))
                for item in transcript]

    @staticmethod
    def _caption_transcript(metadata: dict[str, Any]) -> list[TranscriptSegment]:
        tracks = metadata.get('subtitles') or metadata.get('automatic_captions') or {}
        track = next(iter(tracks.values()), [])
        target = next((item for item in track if item.get('ext') in {'vtt', 'srv3', 'json3'}), None)
        if not target or not target.get('url'):
            return []
        with urllib.request.urlopen(target['url'], timeout=30) as response:
            payload = response.read().decode('utf-8', errors='replace')
        segments = []
        for match in re.finditer(r'(\d{2}:\d{2}:\d{2}[.,]\d{3}) --> .*?\n(.+?)(?=\n\s*\n|\Z)', payload, re.S):
            start = _vtt_time(match.group(1))
            text = re.sub(r'<[^>]+>', '', ' '.join(match.group(2).split()))
            if text:
                segments.append(TranscriptSegment(start, 0.0, text))
        return segments

    def _audio_transcript(self, reference: YouTubeReference) -> list[TranscriptSegment]:
        """Download audio with yt-dlp and transcribe it through Gemini."""
        try:
            import yt_dlp
            from google.genai import types
        except ImportError as exc:
            raise RuntimeError('yt-dlp and google-genai are required for audio transcription') from exc
        with tempfile.TemporaryDirectory(prefix='aura-youtube-') as folder:
            target = Path(folder) / 'audio.%(ext)s'
            options = {
                'quiet': True, 'no_warnings': True, 'noplaylist': True,
                'format': 'bestaudio/best', 'outtmpl': str(target),
            }
            with yt_dlp.YoutubeDL(options) as downloader:
                info = downloader.extract_info(reference.canonical_url, download=True)
            files = list(Path(folder).glob('audio.*'))
            if not files:
                raise RuntimeError('yt-dlp downloaded no audio file')
            audio = files[0]
            mime = str(info.get('ext') or 'webm')
            response = gemini.text([
                types.Part.from_bytes(data=audio.read_bytes(), mime_type=f'audio/{mime}'),
                'Transcribe this audio accurately. Return only the spoken words.',
            ], tier=gemini.SMART, timeout_ms=180_000)
            if not response:
                raise RuntimeError('audio transcription returned no text')
            return [TranscriptSegment(0.0, float(info.get('duration') or 0), response)]

    def transcript(self, reference: YouTubeReference, metadata: dict[str, Any]) -> tuple[list[TranscriptSegment], str]:
        """Try official captions, yt-dlp captions, then audio transcription."""
        try:
            official = self._official_transcript(reference.video_id)
            if official:
                return official, 'official captions'
        except Exception as exc:
            self.notify(f'Transcript fallback: official captions unavailable ({exc})')
        try:
            captions = self._caption_transcript(metadata)
            if captions:
                return captions, 'yt-dlp captions'
        except Exception as exc:
            self.notify(f'Transcript fallback: subtitle download failed ({exc})')
        try:
            audio = self._audio_transcript(reference)
            if audio:
                return audio, 'Gemini audio transcription'
        except Exception as exc:
            self.notify(f'Transcript fallback: audio transcription failed ({exc})')
        raise RuntimeError('no captions were available and audio transcription failed')

    def summarize(self, reference: YouTubeReference, metadata: dict[str, Any],
                  segments: list[TranscriptSegment]) -> str:
        """Map-reduce long transcripts and render a structured Markdown note."""
        chunks = chunk_transcript(segments)
        partials = []
        for index, chunk in enumerate(chunks, 1):
            self.notify(f'Summarizing transcript chunk {index}/{len(chunks)}')
            text = '\n'.join(f'[{_timestamp_link(reference, item.start)}] {item.text}' for item in chunk)
            response = gemini.text(
                'Summarize this transcript chunk with factual key points and definitions.\n' + text,
                tier=gemini.SMART, timeout_ms=90_000,
            )
            if not response:
                raise RuntimeError(f'Gemini returned no summary for chunk {index}')
            partials.append(response)
        merged = gemini.text(
            'Merge these transcript summaries into a factual structured note. Include TL;DR, key points, definitions, action items, and short quotes.\n' +
            '\n\n'.join(partials), tier=gemini.SMART, timeout_ms=90_000,
        )
        if not merged:
            raise RuntimeError('Gemini returned no final note')
        detail = '\n'.join(
            f'- [{_format_time(item.start)}]({_timestamp_link(reference, item.start)}) {item.text}'
            for item in segments
        )
        duration = _format_time(metadata.get('duration', 0))
        return (
            f"# {metadata['title']}\n\n"
            f"- **Channel:** {metadata['channel']}\n"
            f"- **Published:** {metadata.get('publish_date') or 'Unknown'}\n"
            f"- **Duration:** {duration}\n"
            f"- **Transcript source:** {metadata.get('transcript_source', 'captions')}\n"
            f"- **Source:** [{reference.canonical_url}]({reference.canonical_url})\n\n"
            f"## TL;DR\n\n{merged}\n\n"
            f"## Detailed Notes\n\n{detail}\n\n"
            f"## Source\n\n[{reference.canonical_url}]({reference.canonical_url})\n"
        )

    async def run(self, url: str, export: str | None = None) -> Path:
        """Run metadata, transcript, summarization, and atomic note save."""
        reference = parse_youtube_url(url)
        context = WorkflowContext(inputs={'reference': reference})

        def fetch_metadata(_context):
            return self.metadata(reference)

        def fetch_transcript(ctx):
            metadata = ctx.values['metadata']
            transcript, source = self.transcript(reference, metadata)
            metadata['transcript_source'] = source
            return transcript

        def make_note(ctx):
            return self.summarize(reference, ctx.values['metadata'], ctx.values['transcript'])

        def save_note(ctx):
            metadata = ctx.values['metadata']
            self.notes_dir.mkdir(parents=True, exist_ok=True)
            target = self.notes_dir / f"{_safe_filename(metadata['title'])}.md"
            temporary = target.with_suffix('.md.tmp')
            temporary.write_text(ctx.values['note'], encoding='utf-8')
            temporary.replace(target)
            if not target.is_file() or target.stat().st_size == 0:
                raise RuntimeError('saved note failed verification')
            if export:
                _export_note(target, ctx.values['note'], export)
            return target

        runner = WorkflowRunner(self.notify)
        await runner.run([
            WorkflowStep('metadata', fetch_metadata, lambda value, _: bool(value.get('title'))),
            WorkflowStep('transcript', fetch_transcript, lambda value, _: bool(value)),
            WorkflowStep('note', make_note, lambda value, _: bool(value.strip())),
            WorkflowStep('save', save_note, lambda value, _: isinstance(value, Path)),
        ], context)
        return context.values['save']


def _vtt_time(value: str) -> float:
    parts = value.replace(',', '.').split(':')
    return float(parts[0]) * 3600 + float(parts[1]) * 60 + float(parts[2])


def _format_time(seconds: float | int) -> str:
    total = max(0, int(seconds or 0))
    return f'{total // 3600:02d}:{(total % 3600) // 60:02d}:{total % 60:02d}'


def _export_note(markdown_path: Path, content: str, export: str) -> Path:
    """Export Markdown notes to DOCX or PDF and verify the result."""
    kind = export.casefold().lstrip('.')
    target = markdown_path.with_suffix('.docx' if kind == 'docx' else '.pdf')
    if kind == 'docx':
        try:
            from docx import Document
        except ImportError as exc:
            raise RuntimeError('python-docx is required for DOCX export') from exc
        document = Document()
        for line in content.splitlines():
            if line.startswith('# '):
                document.add_heading(line[2:], 0)
            elif line.startswith('## '):
                document.add_heading(line[3:], 1)
            elif line.strip():
                document.add_paragraph(line.lstrip('- '))
        document.save(target)
    elif kind == 'pdf':
        try:
            from reportlab.lib.pagesizes import letter
            from reportlab.lib.styles import getSampleStyleSheet
            from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
        except ImportError as exc:
            raise RuntimeError('reportlab is required for PDF export') from exc
        styles = getSampleStyleSheet()
        story = []
        for line in content.splitlines():
            if not line.strip():
                story.append(Spacer(1, 8))
            else:
                story.append(Paragraph(line.replace('&', '&amp;'), styles['BodyText']))
        SimpleDocTemplate(str(target), pagesize=letter).build(story)
    else:
        raise ValueError('export must be docx or pdf')
    if not target.is_file() or target.stat().st_size == 0:
        raise RuntimeError(f'{kind} export failed verification')
    return target